#!/usr/bin/env python3
"""Build an ATS-friendly .docx from a résumé page on this site.

Applicant tracking systems read Word files more reliably than PDFs, and a
recruiter uploading to a client portal often has to send one. The file is
generated from the HTML page so the two never say different things:

    pip install -r scripts/requirements-resume.txt
    python3 scripts/resume-docx.py resume-it-support.html \
        assets/resume/matthew-mccluster-resume-it-support.docx

Deliberately plain: one column, real Word headings, no tables, text boxes,
images or columns, standard fonts. (The npm package called "docx" in
package.json is an unrelated JavaScript library.)
scripts/test/resume-docx-sync.test.mjs fails if the committed file no
longer carries everything the page says, so re-run this after editing the
page.
"""
import re
import sys
from html import unescape
from html.parser import HTMLParser

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


class Resume(HTMLParser):
    """Collects the bio page's parts in document order."""

    def __init__(self):
        super().__init__()
        self.stack = []
        self.buf = None
        self.name = self.role = self.contact = self.lede = ""
        self.sections = []  # [title, [items]]
        self.job = None

    def _cls(self, attrs):
        return dict(attrs).get("class", "") or ""

    VOID = {"br", "img", "meta", "link", "input", "hr", "source", "wbr"}

    def handle_starttag(self, tag, attrs):
        if tag in self.VOID:  # no end tag: never on the stack
            return
        cls = self._cls(attrs)
        self.stack.append((tag, cls))
        if tag in ("script", "style"):
            return
        if tag == "h1" or tag == "h2" or "bio__role" in cls or "rsm-contact" in cls or "bio__lede" in cls \
                or cls in ("yr", "ti", "wh"):
            self.buf = []
        elif tag == "li" and self._in("bio__facts"):
            self.buf = []
        elif tag == "li" and self._in("dt"):
            self.buf = []
        elif tag == "li" and self._in("tl"):
            self.job = {"yr": "", "ti": "", "wh": "", "dt": []}

    def _in(self, cls):
        return any(cls in c.split() for _, c in self.stack[:-1])

    def handle_endtag(self, tag):
        if not self.stack:
            return
        open_tag, cls = self.stack.pop()
        text = None
        if self.buf is not None and open_tag in ("h1", "h2", "p", "span", "li"):
            text = re.sub(r"\s+", " ", unescape("".join(self.buf))).strip()
        if open_tag == "h1":
            self.name, self.buf = text, None
        elif open_tag == "h2":
            self.sections.append([text, []])
            self.buf = None
        elif "bio__role" in cls:
            self.role, self.buf = text, None
        elif "rsm-contact" in cls:
            self.contact, self.buf = text, None
        elif "bio__lede" in cls:
            self.lede, self.buf = text, None
        elif cls in ("yr", "ti", "wh") and self.job is not None:
            self.job[cls], self.buf = text, None
        elif open_tag == "li" and "dt" in [c for _, c in self.stack] and self.job is not None and text is not None:
            self.job["dt"].append(text)
            self.buf = None
        elif open_tag == "li" and self.job is not None and "tl" in [c for _, c in self.stack]:
            if self.sections:
                self.sections[-1][1].append(("job", self.job))
            self.job = None
        elif open_tag == "li" and text is not None and self.sections:
            self.sections[-1][1].append(("fact", text))
            self.buf = None

    def handle_data(self, data):
        if self.buf is not None and not any(t in ("script", "style") for t, _ in self.stack):
            self.buf.append(data)


def build(src, out):
    page = Resume()
    page.feed(open(src, encoding="utf-8").read())
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Pt(54)

    title = doc.add_paragraph()
    run = title.add_run(page.name)
    run.bold = True
    run.font.size = Pt(20)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if page.role:
        doc.add_paragraph(page.role).runs[0].bold = True
    if page.contact:
        doc.add_paragraph(re.sub(r"^\s*·\s*|\s*·\s*$", "", page.contact))

    doc.add_heading("Summary", level=2)
    doc.add_paragraph(page.lede)

    for heading, items in page.sections:
        if not items:
            continue
        doc.add_heading(heading, level=2)
        for kind, item in items:
            if kind == "fact":
                label, _, rest = item.partition(":")
                p = doc.add_paragraph(style="List Bullet")
                if rest:
                    p.add_run(label + ":").bold = True
                    p.add_run(rest)
                else:
                    p.add_run(item)
            else:
                p = doc.add_paragraph()
                p.add_run(item["ti"]).bold = True
                p.add_run(" · " + item["wh"])
                p.add_run("\n" + re.sub(r"\s+to\s+", " – ", item["yr"])).italic = True
                for line in item["dt"]:
                    doc.add_paragraph(line, style="List Bullet")
    doc.core_properties.author = page.name
    doc.core_properties.title = f"{page.name} résumé"
    doc.save(out)
    return page


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("usage: resume-docx.py <page.html> <out.docx>")
    page = build(sys.argv[1], sys.argv[2])
    print(f"wrote {sys.argv[2]}: {len(page.sections)} sections")
